"""Tests for the export/receipt slice (app/export.py + /jobs/{id}/export.*).
Every export carries only the transcribed text; the receipt is its own endpoint.

Run from fastapi/sqlite/:  python -m pytest tests/test_export.py -v
Works offline: the text-layer test feeds a hand-made page JSON.
"""

import hashlib

import cv2
import numpy as np
import pytest
from fastapi.testclient import TestClient

from app import db, export, storage
from app.main import app

TAMIL_WORD = "வாலறிவன்"


@pytest.fixture(scope="module")
def client(tmp_path_factory):
    tmp = tmp_path_factory.mktemp("exportrun")
    db.DB_PATH = tmp / "test.db"
    storage.UPLOAD_ROOT = tmp / "uploads"
    db.init_db()
    with TestClient(app) as c:
        yield c


@pytest.fixture(scope="module")
def job_id(client, tmp_path_factory):
    img = np.full((800, 1200, 3), 235, np.uint8)
    cv2.putText(img, "TAMIL TEXT LINE", (100, 200), cv2.FONT_HERSHEY_SIMPLEX, 1.5, (20, 20, 20), 3)
    ok, buf = cv2.imencode(".png", img)
    r = client.post("/jobs", files={"file": ("page.png", buf.tobytes(), "image/png")})
    assert r.status_code == 200
    return r.json()["job_id"]


def _page(lines):
    return {"page": 1, "profile": "FAST",
            "quality": {"blur": 1, "contrast": 1, "noise": 1, "skew_deg": 0},
            "lines": lines, "text": "\n".join(l["body"] for l in lines),
            "needs_review": False, "processing_ms": 10}


def test_pdf_text_layer_has_tamil_word(tmp_path):
    import pypdfium2 as pdfium
    img = np.full((600, 1200, 3), 240, np.uint8)
    master = tmp_path / "master.png"
    cv2.imwrite(str(master), img)
    page = _page([{"id": "L1", "seq": 1, "body": "கற்றதனால் ஆய பயனென்கொல் " + TAMIL_WORD,
                   "bbox": [100, 100, 1000, 80], "confidence": 0.9}])
    job = {"id": "x" * 32, "filename": "t.png", "sha256": "0" * 64,
           "status": "done", "created_at": "2026-01-01T00:00:00+00:00"}
    receipt = export.build_receipt(job, [page], reviewer="R", ocr_engine="sarvam")
    data = export.build_pdf(master, [page], receipt)
    assert data.startswith(b"%PDF")
    pdf = pdfium.PdfDocument(data)
    assert len(pdf) == 2  # visible text page + scan page, no receipt page
    assert TAMIL_WORD in pdf[0].get_textpage().get_text_range()  # visible text
    assert TAMIL_WORD in pdf[1].get_textpage().get_text_range()  # scan layer
    for i in range(len(pdf)):
        assert "0" * 64 not in pdf[i].get_textpage().get_text_range()
    assert pdf.get_metadata_dict()["Keywords"].startswith("master-sha256:" + "0" * 64)


RECEIPT_MARKERS = ("Pink Cloud processing receipt", "processing receipt",
                   "Master SHA-256", "Job:", "Reviewer:", "Processing time")

POEM = ["அகர முதல எழுத்தெல்லாம் ஆதி",
        "பகவன் முதற்றே உலகு",
        "கற்றதனால் ஆய பயனென்கொல் " + TAMIL_WORD]


def _real_scan_job(client, lines, name="50823516_poem_44.jpg"):
    """Upload a real image, then store known OCR lines as the job result."""
    import json
    img = np.full((1200, 1600, 3), 235, np.uint8)
    cv2.putText(img, "SCAN", (200, 300), cv2.FONT_HERSHEY_SIMPLEX, 4, (20, 20, 20), 8)
    ok, buf = cv2.imencode(".jpg", img)
    jid = client.post("/jobs", files={"file": (name, buf.tobytes(), "image/jpeg")}).json()["job_id"]
    page = _page([{"id": f"L{i}", "seq": i, "body": b,
                   "bbox": [100, 100 + i * 70, 1200, 60], "confidence": 0.95}
                  for i, b in enumerate(lines, 1)])
    db.set_result(jid, "done", json.dumps({"pages": [page]}, ensure_ascii=False))
    return jid


def _ink(pdf_page) -> float:
    """Share of non-white pixels when the page is rendered."""
    g = np.asarray(pdf_page.render(scale=1).to_pil().convert("L"))
    return float((g < 200).mean())


def test_export_pdf_page1_is_visible_ocr_text_no_receipt(client):
    """Regression (user report): the PDF must open with his document's
    recognized text, visibly, and carry no receipt page at all."""
    import pypdfium2 as pdfium
    jid = _real_scan_job(client, POEM)
    r = client.get(f"/jobs/{jid}/export.pdf")
    assert r.status_code == 200
    pdf = pdfium.PdfDocument(r.content)
    assert len(pdf) == 2  # text page, scan page
    first = pdf[0].get_textpage().get_text_range()
    for line in POEM:
        assert line in first, line
    assert _ink(pdf[0]) > 0.003  # glyphs are actually drawn, not an invisible layer
    # the visible text is real vector glyphs (paths), the scan is page 2
    kinds0 = {o.type for o in pdf[0].get_objects()}
    assert pdfium.raw.FPDF_PAGEOBJ_PATH in kinds0
    assert pdfium.raw.FPDF_PAGEOBJ_IMAGE not in kinds0
    assert pdfium.raw.FPDF_PAGEOBJ_IMAGE in {o.type for o in pdf[1].get_objects()}
    assert all(line in pdf[1].get_textpage().get_text_range() for line in POEM)
    for i in range(len(pdf)):
        text = pdf[i].get_textpage().get_text_range()
        for marker in RECEIPT_MARKERS:
            assert marker not in text, (i, marker)
    assert "PinkCloudReceipt" not in pdf.get_metadata_dict()
    # old links with ?receipt_page=true still get no receipt page
    old = pdfium.PdfDocument(client.get(f"/jobs/{jid}/export.pdf",
                                        params={"receipt_page": True}).content)
    assert len(old) == 2
    # the receipt data itself is still served by its own endpoint
    rc = client.get(f"/jobs/{jid}/receipt")
    assert rc.status_code == 200 and rc.json()["lines"]["total"] == len(POEM)


def test_export_pdf_long_text_is_fitted_onto_its_one_text_page(client):
    """A long page used to flow over several A4 text pages; now each source
    page's text is shrunk/reflowed to fit ONE text page (in reading order)."""
    import pypdfium2 as pdfium
    lines = [f"{i:03d} " + POEM[i % 3] for i in range(90)]
    jid = _real_scan_job(client, lines, name="long.jpg")
    pdf = pdfium.PdfDocument(client.get(f"/jobs/{jid}/export.pdf").content)
    assert len(pdf) == 2  # one fitted text page + the scan
    visible = pdf[0].get_textpage().get_text_range()
    pos = [visible.index(l) for l in lines]
    assert pos == sorted(pos)  # reading order kept
    assert pdfium.raw.FPDF_PAGEOBJ_IMAGE in {o.type for o in pdf[1].get_objects()}
    # every drawn glyph sits inside the page margins (nothing overflows)
    w, h = pdf[0].get_size()
    for o in pdf[0].get_objects():
        l, b, r, t = o.get_bounds()
        assert l >= export.TEXT_MARGIN - 1 and r <= w - export.TEXT_MARGIN + 1, (l, r)
        assert b >= export.TEXT_MARGIN * 0.5 and t <= h - export.TEXT_MARGIN * 0.5, (b, t)


def _fit_pdf(tmp_path, page_lines, shape=(1600, 1200)):
    """build_pdf over len(page_lines) blank scans of `shape` (h, w)."""
    import pypdfium2 as pdfium
    masters, pages = [], []
    for n, lines in enumerate(page_lines, 1):
        m = tmp_path / f"m{n}.png"
        cv2.imwrite(str(m), np.full((shape[0], shape[1], 3), 240, np.uint8))
        masters.append(m)
        p = _page([{"id": f"L{i}", "seq": i, "body": b,
                    "bbox": [50, 50 + i * 20, 900, 18], "confidence": 0.95}
                   for i, b in enumerate(lines, 1)])
        p["page"] = n
        pages.append(p)
    job = {"id": "f" * 32, "filename": "fit.pdf", "sha256": "0" * 64,
           "status": "done", "created_at": "2026-01-01T00:00:00+00:00"}
    receipt = export.build_receipt(job, pages)
    return pdfium.PdfDocument(export.build_pdf(masters, pages, receipt))


PARA = ("அகர முதல எழுத்தெல்லாம் ஆதி பகவன் முதற்றே உலகு கற்றதனால் ஆய "
        "பயனென்கொல் வாலறிவன் நற்றாள் தொழாஅர் எனின் மலர்மிசை ஏகினான் "
        "மாணடி சேர்ந்தார் நிலமிசை நீடுவாழ் வார்")


def test_fit_text_to_page_shrinks_and_wraps():
    font_data = export.FONT_PATH.read_bytes()
    hb, upem = export._hb_font(font_data)
    # short text keeps the natural size
    size, vis = export.fit_text_to_page(hb, upem, POEM, 483, 730)
    assert size == export.TEXT_SIZE and vis == POEM
    # a long paragraph wraps inside the width
    size, vis = export.fit_text_to_page(hb, upem, [PARA], 200, 730)
    assert len(vis) > 1
    assert all(export._shaped_width(hb, upem, v, size) <= 200.01 for v in vis)
    assert " ".join(vis) == PARA
    # lots of paragraphs: the font shrinks until the block fits the height
    many = [PARA] * 12
    size, vis = export.fit_text_to_page(hb, upem, many, 483, 730)
    assert export.FIT_MIN_SIZE <= size < export.TEXT_SIZE
    need = size + (len(vis) - 1) * export.leading_for(size) + 0.3 * size
    assert need <= 730


def test_pdf_one_fitted_text_page_per_source_page(tmp_path):
    """Multi-page job: text page N carries exactly source page N's text
    (labelled "Page N"), then the N scans follow."""
    dense = [PARA] * 10
    pdf = _fit_pdf(tmp_path, [POEM, dense, []])
    assert len(pdf) == 6  # 3 text pages + 3 scans
    t = [pdf[i].get_textpage().get_text_range() for i in range(3)]
    assert "Page 1" in t[0] and POEM[0] in t[0] and "மலர்மிசை" not in t[0]
    assert t[0].count("வாலறிவன்") == 1
    assert "Page 2" in t[1] and "Page 1" not in t[1] and t[1].count("வாலறிவன்") == 10
    assert "Page 3" in t[2] and export.EMPTY_PAGE_NOTE in t[2]
    for i in range(3, 6):
        assert pdfium_image_page(pdf[i])


def pdfium_image_page(page) -> bool:
    import pypdfium2 as pdfium
    return pdfium.raw.FPDF_PAGEOBJ_IMAGE in {o.type for o in page.get_objects()}


def test_pdf_text_pages_are_a4_portrait_for_any_scan(tmp_path):
    for shape in ((1000, 1600), (1600, 1000)):
        w, h = _fit_pdf(tmp_path, [POEM], shape=shape)[0].get_size()
        assert (round(w), round(h)) == (595, 842)


def test_pdf_extreme_page_continues_instead_of_clipping(tmp_path):
    """Too dense even at FIT_MIN_SIZE: continue on a "(continued)" page,
    never drop text."""
    lines = [f"{i:04d} {PARA}" for i in range(120)]
    pdf = _fit_pdf(tmp_path, [lines])
    assert len(pdf) >= 3  # >=2 text pages + scan
    n_text = len(pdf) - 1
    assert "(continued)" in pdf[1].get_textpage().get_text_range()
    visible = "\n".join(pdf[i].get_textpage().get_text_range() for i in range(n_text))
    pos = [visible.index(f"{i:04d}") for i in range(120)]
    assert pos == sorted(pos)
    assert pdfium_image_page(pdf[len(pdf) - 1])


def test_receipt_counts(client, job_id):
    r = client.get(f"/jobs/{job_id}/receipt", params={"reviewer": "Anu"})
    assert r.status_code == 200
    rc = r.json()
    assert rc["page_count"] == 1
    assert rc["reviewer"] == "Anu"
    assert rc["lines"]["auto"] + rc["lines"]["human_review"] == rc["lines"]["total"]
    assert rc["pages"]["auto"] + rc["pages"]["human_review"] == 1
    master = storage.master_path(job_id)
    assert rc["master"]["sha256"] == hashlib.sha256(master.read_bytes()).hexdigest()
    assert rc["master"]["verified_on_disk"] is True


def test_export_endpoints(client, job_id):
    p = client.get(f"/jobs/{job_id}/export.pdf")
    assert p.status_code == 200 and p.content.startswith(b"%PDF")
    assert p.headers["content-type"] == "application/pdf"
    t = client.get(f"/jobs/{job_id}/export.txt")
    assert t.status_code == 200
    # transcribed text only: no "# " header, no receipt, no page separator
    # on a single-page job
    assert not t.text.startswith("#") and "=== page" not in t.text
    for marker in RECEIPT_MARKERS:
        assert marker not in t.text


def test_unknown_job_404(client):
    for path in ("receipt", "export.pdf", "export.txt", "export.docx"):
        assert client.get(f"/jobs/{'f' * 32}/{path}").status_code == 404


def test_tamil_filename_export_no_500(client):
    """Non-ASCII upload names must not crash the latin-1 download headers."""
    img = np.full((400, 600, 3), 235, np.uint8)
    ok, buf = cv2.imencode(".png", img)
    r = client.post("/jobs", files={"file": ("தமிழ்.png", buf.tobytes(), "image/png")})
    assert r.status_code == 200
    jid = r.json()["job_id"]
    for path in ("export.txt", "export.pdf", "export.docx"):
        resp = client.get(f"/jobs/{jid}/{path}")
        assert resp.status_code == 200
        cd = resp.headers["content-disposition"]
        assert 'filename="pinkcloud.' in cd and "filename*=UTF-8''" in cd


def test_saved_corrections_applied_to_exports(client):
    import json
    import pypdfium2 as pdfium
    img = np.full((600, 1200, 3), 240, np.uint8)
    # distinct bytes per test: identical uploads dedup to the existing job
    cv2.putText(img, "EXPORT-C", (20, 40), cv2.FONT_HERSHEY_SIMPLEX, 1.0, (20, 20, 20), 2)
    ok, buf = cv2.imencode(".png", img)
    jid = client.post("/jobs", files={"file": ("c.png", buf.tobytes(), "image/png")}).json()["job_id"]
    raw, fixed = "வாழறிவன்", TAMIL_WORD
    page = _page([{"id": "L1", "seq": 1, "body": "கற்றதனால் ஆய " + raw,
                   "bbox": [100, 100, 1000, 80], "confidence": 0.9}])
    db.set_result(jid, "done", json.dumps({"pages": [page]}, ensure_ascii=False))
    cpath = storage.UPLOAD_ROOT / jid / "corrections.json"

    cpath.write_text("{not json", encoding="utf-8")  # malformed -> raw OCR
    assert raw in client.get(f"/jobs/{jid}/export.txt").text

    cpath.write_text(json.dumps({"corrections": [
        {"page": 1, "line": "L1", "word": 3, "before": raw, "after": fixed},
        {"page": 1, "line": "L1", "word": 1, "before": "stale", "after": "X"},  # skipped
    ]}, ensure_ascii=False), encoding="utf-8")
    txt = client.get(f"/jobs/{jid}/export.txt").text
    assert fixed in txt and raw not in txt and "கற்றதனால்" in txt
    pdf = pdfium.PdfDocument(client.get(f"/jobs/{jid}/export.pdf").content)
    layer = pdf[0].get_textpage().get_text_range()
    assert fixed in layer and raw not in layer
    rc = client.get(f"/jobs/{jid}/receipt").json()
    assert rc["corrections"]["total"] == 1 and rc["corrections"]["by_tier"] == {"human": 1}


# --------------------------------------------------------------------------
# DOCX export
# --------------------------------------------------------------------------

DOCX_TYPE = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def _docx_text(data: bytes) -> str:
    import io
    import zipfile
    from docx import Document
    assert zipfile.is_zipfile(io.BytesIO(data))
    assert "word/document.xml" in zipfile.ZipFile(io.BytesIO(data)).namelist()
    return "\n".join(p.text for p in Document(io.BytesIO(data)).paragraphs)


def test_export_docx_endpoint(client, job_id):
    r = client.get(f"/jobs/{job_id}/export.docx", params={"reviewer": "Anu"})
    assert r.status_code == 200
    assert r.headers["content-type"] == DOCX_TYPE
    assert r.content[:2] == b"PK"
    assert 'filename="page.docx"' in r.headers["content-disposition"]
    sha = hashlib.sha256(storage.master_path(job_id).read_bytes()).hexdigest()
    assert r.headers["x-master-sha256"] == sha
    text = _docx_text(r.content)
    # transcribed text only: no title block, no receipt section
    assert "Pink Cloud export" not in text and "Page 1" not in text
    assert sha not in text and "Anu" not in text
    for marker in RECEIPT_MARKERS:
        assert marker not in text
    # provenance stays in the document properties only (like the PDF Info)
    from docx import Document
    import io
    props = Document(io.BytesIO(r.content)).core_properties
    assert props.keywords == "master-sha256:" + sha


def test_export_docx_corrections_applied(client):
    import json
    img = np.full((600, 1200, 3), 240, np.uint8)
    cv2.putText(img, "EXPORT-D", (20, 40), cv2.FONT_HERSHEY_SIMPLEX, 1.0, (20, 20, 20), 2)
    ok, buf = cv2.imencode(".png", img)
    jid = client.post("/jobs", files={"file": ("d.png", buf.tobytes(), "image/png")}).json()["job_id"]
    raw, fixed = "வாழறிவன்", TAMIL_WORD
    page = _page([{"id": "L1", "seq": 1, "body": "கற்றதனால் ஆய " + raw,
                   "bbox": [100, 100, 1000, 80], "confidence": 0.9}])
    db.set_result(jid, "done", json.dumps({"pages": [page]}, ensure_ascii=False))
    assert raw in _docx_text(client.get(f"/jobs/{jid}/export.docx").content)
    (storage.UPLOAD_ROOT / jid / "corrections.json").write_text(json.dumps({"corrections": [
        {"page": 1, "line": "L1", "word": 3, "before": raw, "after": fixed},
    ]}, ensure_ascii=False), encoding="utf-8")
    text = _docx_text(client.get(f"/jobs/{jid}/export.docx").content)
    assert fixed in text and raw not in text and "கற்றதனால்" in text
    assert "Corrections" not in text and "receipt" not in text.lower()


def test_export_docx_unfinished_job_409(client):
    img = np.full((400, 600, 3), 235, np.uint8)
    cv2.putText(img, "EXPORT-Q", (20, 40), cv2.FONT_HERSHEY_SIMPLEX, 1.0, (20, 20, 20), 2)
    ok, buf = cv2.imencode(".png", img)
    jid = client.post("/jobs", files={"file": ("q.png", buf.tobytes(), "image/png")}).json()["job_id"]
    db.set_result(jid, "processing", None)
    assert client.get(f"/jobs/{jid}/export.docx").status_code == 409


def test_export_docx_unknown_job_404(client):
    assert client.get(f"/jobs/{'e' * 32}/export.docx").status_code == 404


# --------------------------------------------------------------------------
# every export carries ONLY the transcribed text (no receipt anywhere)
# --------------------------------------------------------------------------

def test_txt_export_is_only_transcribed_text(client):
    jid = _real_scan_job(client, POEM, name="poem.jpg")
    txt = client.get(f"/jobs/{jid}/export.txt", params={"reviewer": "Anu"}).text
    assert txt == "\n".join(POEM) + "\n"
    rc = client.get(f"/jobs/{jid}/receipt")  # receipt endpoint stays
    assert rc.status_code == 200 and rc.json()["reviewer"] is None


def test_docx_export_is_only_transcribed_text(client):
    import io
    from docx import Document
    jid = _real_scan_job(client, POEM, name="poem2.jpg")
    doc = Document(io.BytesIO(client.get(f"/jobs/{jid}/export.docx").content))
    assert [p.text for p in doc.paragraphs] == POEM


def test_multi_page_txt_and_docx_page_separators():
    import io
    from docx import Document
    p1 = _page([{"id": "L1", "seq": 1, "body": POEM[0], "bbox": [0, 0, 1, 1], "confidence": 0.9}])
    p2 = dict(_page([{"id": "L1", "seq": 1, "body": POEM[1], "bbox": [0, 0, 1, 1], "confidence": 0.9}]), page=2)
    txt = export.build_txt([p2, p1])
    assert txt == f"=== page 1 ===\n{POEM[0]}\n\n=== page 2 ===\n{POEM[1]}\n"
    doc = Document(io.BytesIO(export.build_docx([p1, p2])))
    assert [p.text for p in doc.paragraphs] == ["Page 1", POEM[0], "Page 2", POEM[1]]


def test_empty_job_exports():
    import io
    from docx import Document
    p = _page([])
    p["text"] = ""
    assert export.build_txt([p]) == ""
    doc = Document(io.BytesIO(export.build_docx([p])))
    assert [x.text for x in doc.paragraphs] == [export.EMPTY_TEXT_NOTE]


# ---- "Include original scans" toggle (export step) ------------------------

def _scan_pages(pdf) -> list[int]:
    import pypdfium2 as pdfium
    return [i for i in range(len(pdf))
            if pdfium.raw.FPDF_PAGEOBJ_IMAGE in {o.type for o in pdf[i].get_objects()}]


def test_build_pdf_include_scans_default_and_on_keep_scan_pages(tmp_path):
    """Default (and include_scans=True) is the old output: N text pages,
    then N scans, each scan with its invisible Tamil layer."""
    dense = [PARA] * 10
    pdf = _fit_pdf(tmp_path, [POEM, dense])
    assert len(pdf) == 4 and _scan_pages(pdf) == [2, 3]
    assert all(l in pdf[2].get_textpage().get_text_range() for l in POEM)
    assert "scan" in pdf.get_metadata_dict()["Subject"]


def test_build_pdf_include_scans_off_is_text_pages_only(tmp_path):
    """include_scans=False: only the fitted text pages (one per source page),
    no images, and the text is still extractable/searchable."""
    import pypdfium2 as pdfium
    dense = [PARA] * 10
    masters, pages = [], []
    for n, lines in enumerate([POEM, dense], 1):
        m = tmp_path / f"m{n}.png"
        cv2.imwrite(str(m), np.full((1600, 1200, 3), 240, np.uint8))
        masters.append(m)
        p = _page([{"id": f"L{i}", "seq": i, "body": b,
                    "bbox": [50, 50 + i * 20, 900, 18], "confidence": 0.95}
                   for i, b in enumerate(lines, 1)])
        p["page"] = n
        pages.append(p)
    job = {"id": "f" * 32, "filename": "fit.pdf", "sha256": "0" * 64,
           "status": "done", "created_at": "2026-01-01T00:00:00+00:00"}
    receipt = export.build_receipt(job, pages)
    on = pdfium.PdfDocument(export.build_pdf(masters, pages, receipt, include_scans=True))
    off = pdfium.PdfDocument(export.build_pdf(masters, pages, receipt, include_scans=False))
    assert len(on) == 4
    assert len(off) == 2 and _scan_pages(off) == []
    t = [off[i].get_textpage().get_text_range() for i in range(2)]
    assert "Page 1" in t[0] and all(l in t[0] for l in POEM)
    assert "Page 2" in t[1] and "மலர்மிசை" in t[1] and POEM[2] not in t[1]
    # the text pages are identical in both modes (same fit, same size)
    for i in range(2):
        assert off[i].get_size() == on[i].get_size()
        assert off[i].get_textpage().get_text_range() == on[i].get_textpage().get_text_range()
    assert _ink(off[0]) > 0.003
    meta = off.get_metadata_dict()
    assert meta["Keywords"].startswith("master-sha256:" + "0" * 64)
    assert "not included" in meta["Subject"]


def test_export_pdf_endpoint_include_scans_param(client):
    """GET /export.pdf: missing include_scans = old behaviour (text + scan);
    include_scans=1/true keeps the scan; include_scans=0/false drops it."""
    import pypdfium2 as pdfium
    jid = _real_scan_job(client, POEM, name="toggle.jpg")

    def get(params=None):
        r = client.get(f"/jobs/{jid}/export.pdf", params=params or {})
        assert r.status_code == 200, r.text
        assert r.headers["content-type"] == "application/pdf"
        return pdfium.PdfDocument(r.content)

    for params in (None, {"include_scans": "1"}, {"include_scans": "true"}):
        pdf = get(params)
        assert len(pdf) == 2 and _scan_pages(pdf) == [1], params
    for params in ({"include_scans": "0"}, {"include_scans": "false"}):
        pdf = get(params)
        assert len(pdf) == 1 and _scan_pages(pdf) == [], params
        assert all(l in pdf[0].get_textpage().get_text_range() for l in POEM)
    # still combines with the legacy receipt_page flag
    assert len(get({"include_scans": "0", "receipt_page": "true"})) == 1
    # junk values are rejected, not silently treated as on/off
    assert client.get(f"/jobs/{jid}/export.pdf",
                      params={"include_scans": "maybe"}).status_code == 422
